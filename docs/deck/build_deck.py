"""Build the 5-minute Challenge Day deck as PowerPoint (.pptx) and PDF.

One slide description drives both renderers, so the two files match:
python-pptx writes the editable deck, reportlab draws the PDF copy.
Numbers are read from outputs/public/numbers.json, like the report.

    python docs/deck/build_deck.py
"""
from __future__ import annotations

import json
import re
from pathlib import Path

import yaml
from PIL import Image as PILImage
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas as rl_canvas
from reportlab.platypus import Paragraph, Table, TableStyle

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
FIG = ROOT / "outputs" / "public" / "figures"
N = json.loads((ROOT / "outputs" / "public" / "numbers.json").read_text())
FRONT = yaml.safe_load((ROOT / "docs" / "report" / "underlink_report.md").read_text().split("---", 2)[1])
TEAM = FRONT["team"].replace("[", "").replace("]", "").strip()
STEM = f"DataChallenge_{TEAM}_Slides"

W, H = 13.333, 7.5                      # inches, 16:9
DARK, WHITE, INK, MUTED = "12303B", "FFFFFF", "1A1A1A", "5A6570"
BLUE, ORANGE, TINT, LIGHT = "2A78D6", "EB6834", "EEF4FA", "9CC3F0"
FONT = "Arial"

# ------------------------------------------------------------------ numbers
c, r, rep, fb = N["chains"], N["relays"], N["repair"], N["fallbacks"]
lo, hi = (round(x * 100) for x in N["sensitivity"]["ge1_spof_share_range"])
people_ge1 = f"{round(c['people_ge1_spof'], -2):,}"
jan_slow = rep["jan_bands"].get(">14 d", 0)
sweep_lo = min(x["gt14d"] for x in rep["sweep"]); sweep_hi = max(x["gt14d"] for x in rep["sweep"])


# ------------------------------------------------------------------ slide spec
def t(x, y, w, h, text, size=16, bold=False, color=INK, align="left", anchor="top"):
    return ("text", dict(x=x, y=y, w=w, h=h, text=text, size=size, bold=bold, color=color, align=align, anchor=anchor))


def img(path, x, y, w):
    return ("image", dict(path=str(path), x=x, y=y, w=w))


def box(x, y, w, h, fill=TINT):
    return ("box", dict(x=x, y=y, w=w, h=h, fill=fill))


def dot(x, y, d, fill=ORANGE):
    return ("dot", dict(x=x, y=y, d=d, fill=fill))


def line(x1, y1, x2, y2, color=MUTED, width=2):
    return ("line", dict(x1=x1, y1=y1, x2=x2, y2=y2, color=color, width=width))


def table(x, y, w, col_w, rows, size=12):
    return ("table", dict(x=x, y=y, w=w, col_w=col_w, rows=rows, size=size))


def dotted_list(x, y, w, items, size=16, gap=0.72, color=INK, dot_color=ORANGE):
    """Short items, each marked with a relay dot (the deck's one motif)."""
    out = []
    for i, it in enumerate(items):
        yy = y + i * gap
        out += [dot(x, yy + 0.09, 0.16, dot_color), t(x + 0.32, yy, w - 0.32, gap, it, size=size, color=color)]
    return out


def title(text, color=INK):
    return t(0.6, 0.45, W - 1.2, 0.9, text, size=32, bold=True, color=color)


def footer(dark=False):
    col = LIGHT if dark else MUTED
    return t(0.6, H - 0.45, W - 1.2, 0.3, f"{TEAM}  |  Underlink  |  CDU IT Code Fair 2026, Data Innovation Challenge", size=10, color=col)


def stat(x, y, w, big, label, big_color=BLUE):
    return [t(x, y, w, 0.9, big, size=44, bold=True, color=big_color), t(x, y + 0.85, w, 1.0, label, size=14, color=INK)]


def chain_motif(x, y, n=6, gap=0.62, d=0.26):
    """A fibre block, relays in a row, and the community: the Underlink motif."""
    els = [("rect", dict(x=x, y=y - 0.16, w=0.34, h=0.58, fill=BLUE)), line(x + 0.34, y + 0.13, x + 0.34 + gap * (n + 1), y + 0.13, color=LIGHT, width=2)]
    for i in range(n):
        els.append(dot(x + 0.34 + gap * (i + 1) - d / 2, y, d, ORANGE))
    els.append(dot(x + 0.34 + gap * (n + 1) - 0.1, y + 0.03, 0.2, WHITE))
    return els


SLIDES = [
    dict(bg=DARK, els=[
        t(0.6, 0.55, 9, 0.4, "CDU IT Code Fair 2026  |  Data Innovation Challenge: Remote Connectivity", size=13, color=LIGHT),
        t(0.6, 1.7, 11, 1.2, "Underlink", size=60, bold=True, color=WHITE),
        t(0.6, 2.85, 11.5, 0.8, "Revealing the Hidden Dependencies Behind NT Connectivity", size=26, color=LIGHT),
        *chain_motif(0.7, 4.25),
        t(0.6, 5.15, 11, 0.5, f"{TEAM}  |  Harsh Rastogi  |  Aashish", size=18, bold=True, color=WHITE),
        t(0.6, 6.25, 11.5, 0.7, "We acknowledge the Larrakia people, Traditional Owners of the land where we present today. "
                               "This presentation contains no images or voices of people.", size=12, color=LIGHT),
    ], notes="Acknowledge Country. We are Team DIC017, Harsh and Aashish. Underlink looks at what remote NT phone service "
             "depends on, and what happens on the day that breaks. About 30 seconds."),

    dict(bg=WHITE, els=[
        title("Places the map calls covered still go dark"),
        box(0.6, 1.6, 3.9, 2.45), box(4.72, 1.6, 3.9, 2.45), box(8.84, 1.6, 3.9, 2.45),
        t(0.85, 1.8, 3.5, 0.8, "9 weeks", size=36, bold=True, color=ORANGE),
        t(0.85, 2.65, 3.45, 1.8, "Wadeye, February to April 2026. Rolling outages after flooding. People could not reach 000 or top up prepaid power cards.", size=14),
        t(4.97, 1.8, 3.5, 0.8, "Flat batteries", size=36, bold=True, color=ORANGE),
        t(4.97, 2.65, 3.45, 1.8, "Borroloola, March 2024. Solar-powered sites failed under Ex-Tropical Cyclone Megan. Crews could reach them only by helicopter.", size=14),
        t(9.09, 1.8, 3.5, 0.8, "1,000 km", size=36, bold=True, color=ORANGE),
        t(9.09, 2.65, 3.45, 1.8, "June 2026. One fibre break left parts of the NT offline for about a day. Repair crews came from about 1,000 km away.", size=14),
        t(0.6, 4.45, 12, 0.9, "Coverage maps show signal on a normal day. They do not show what that signal depends on, or how long a repair takes.", size=20, bold=True),
        t(0.6, 5.5, 12, 0.6, "Mobile coverage reaches about 5% of NT land (NT Government estimate). Sources: ABC 7 Apr 2026; AAP 2024; NT Independent 2026.", size=12, color=MUTED),
        footer(),
    ], notes="Three recent NT outages, all in places the coverage map shows as covered. Wadeye lost service on and off for nine weeks. "
             "Borroloola's solar sites ran flat under cloud. A single fibre break took out parts of the Territory. The map answers where "
             "signal is. It does not say what the signal hangs off or how long a fix takes. About 35 seconds."),

    dict(bg=WHITE, els=[
        title("We read a public licence register as a network"),
        *dotted_list(0.6, 1.75, 5.6, [
            "ACMA's licence register lists 808 NT point-to-point radio links licensed to Telstra. Merged by site pair, they form 358 sites and 313 links.",
            "A dominator tree finds every single-path relay, meaning a relay that every licensed radio path from a place to fibre must pass through.",
            "Four more measures sit beside it and are never merged into one score: cyclone replay, backup power, repair window and what still works.",
            "All 8 organiser datasets plus 10 more. One command rebuilds every number and figure.",
        ], size=15, gap=1.12),
        img(FIG / "fig2_pipeline.png", 6.45, 1.7, 6.3),
        footer(),
    ], notes="The ACMA register is public and updated daily. Each licensed point-to-point link has coordinates at both ends, so the "
             "links form a network. For each of 782 places we trace the path back to a fibre town. A dominator tree gives the exact "
             "set of single-path relays, the relays with no way around them. We keep four other measures beside that and never blend them into one score. About 40 seconds."),

    dict(bg=WHITE, els=[
        title(f"{c['with_ge1_spof']} of {c['radio_chain_places']} radio-chain places depend on a single-path relay"),
        *stat(0.6, 1.55, 3.9, f"{c['with_ge1_spof']} of {c['radio_chain_places']}", f"larger radio-chain places depend on at least one single-path relay (about {people_ge1} people)"),
        *stat(4.72, 1.55, 3.9, f"{c['ge1_spof_inside_predicted_4g']} of {c['with_ge1_spof']}", "of those places sit inside Telstra's predicted 4G coverage", big_color=ORANGE),
        *stat(8.84, 1.55, 3.9, f"{lo} to {hi}%", "share with a single-path relay across all 9 sensitivity runs"),
        img(FIG / "fig1_hero_chain.png", 1.4, 3.75, 7.4),
        t(9.05, 3.95, 3.7, 2.6, "Ampilatwatja reaches fibre through five single-path relays in a row. Galiwin'ku has the same shape. "
                                "In 2024 a solar Telstra site ran flat and Galiwin'ku lost reception for 12 nights.", size=13),
        footer(),
    ], notes="Of 23 larger places that reach fibre over radio, 18 depend on at least one single-path relay, and 16 of those are inside Telstra's predicted coverage. Across nine sensitivity runs the share stays "
             "between 78 and 84 percent. Hidden fibre or satellite links can only lower the count, so 18 is an upper bound. About 40 seconds."),

    dict(bg=WHITE, els=[
        title("Backup power is not public, and the wet sets repair time"),
        *stat(0.6, 1.55, 4.3, f"{r['published_autonomy']} of {r['spof_relays']}", "single-path relays are in the Mobile Network Hardening Program, so their battery hours are not public", big_color=ORANGE),
        *stat(0.6, 3.35, 4.3, f"{jan_slow} of {rep['radio_chain_places']}", f"radio-chain places take over 14 days to repair in January, set by road access to a relay ({sweep_lo} to {sweep_hi} across sweeps)"),
        *stat(0.6, 5.1, 4.3, f"{round(rep['median_travel_hours'])} hours", "median July drive from a crew base"),
        img(FIG / "fig5_repair_power.png", 5.3, 1.75, 7.45),
        t(5.3, 5.55, 7.4, 0.9, "The wet-season delay is an assumption, swept from 14 to 60 days. The list of places comes from road distance alone.", size=12, color=MUTED),
        footer(),
    ], notes="None of the 68 single-path relays is in the Hardening Program, so nobody outside the carrier knows "
             "their battery hours. In July, 17 of 23 places are back within a day and the rest within three. In January, 17 of 23 wait more than two weeks "
             "because a radio site on their chain sits more than 10 km off a sealed road. The days are our assumption; the list of places comes from road distance. About 40 seconds."),

    dict(bg=WHITE, els=[
        title("Where the model is right, and where it misses"),
        table(0.6, 1.6, 8.3, [2.3, 3.4, 2.6], [
            ["Reported outage", "Underlink, from public radio data", "Result"],
            ["Ampilatwatja, 2024", "Radio chain, 5 single-path relays", "Consistent (planned works)"],
            ["Galiwin'ku, 2024", "5 single-path relays, none in the Hardening Program", "Fits, if the solar site is on its chain"],
            ["Milingimbi, 2024", "Radio chain, no single-path relay", "Missed: which site serves it"],
            ["Borroloola, 2024", "Fibre town", "Missed: battery hours"],
            ["Wadeye, 2026", "No licensed radio site nearby", "Missed: fibre routes"],
            ["June 2026 fibre break", "Not flagged", "Missed: fibre routes"],
        ], size=13),
        box(9.25, 1.6, 3.5, 3.1, fill=DARK),
        t(9.5, 1.85, 3.05, 3.8, "Every miss traces to data carriers hold and do not publish: battery hours, fibre routes and which site serves which place.\n\nThat is our data request.", size=16, color=WHITE),
        footer(),
    ], notes="We tested the model against every outage we could find in the news. It fits Ampilatwatja and Galiwin'ku, including the "
             "cause at Galiwin'ku. It misses four events, and each miss comes from data only carriers hold. We show the misses because "
             "they tell DCDD exactly which data to ask for. About 40 seconds."),

    dict(bg=WHITE, els=[
        title("Three audiences, and it runs with the Wi-Fi off"),
        box(0.6, 1.55, 3.9, 4.55), box(4.72, 1.55, 3.9, 4.55), box(8.84, 1.55, 3.9, 4.55),
        t(0.85, 1.7, 3.5, 0.5, "Communities", size=20, bold=True, color=BLUE),
        img(ROOT / "docs" / "deck" / "assets" / "card.png", 1.55, 2.25, 2.0),
        t(0.85, 5.1, 3.45, 1.0, "Draft card for co-design. Colour, one word and a symbol, plus a panel for paid interpreters.", size=12),
        t(4.97, 1.7, 3.5, 0.5, "Government (DCDD)", size=20, bold=True, color=BLUE),
        *dotted_list(4.97, 2.35, 3.45, [
            f"{c['with_ge3_spof']} places depend on 3 or more single-path relays",
            f"{r['spof_relays']} of {r['spof_relays']} single-path relays have no public battery hours",
            f"{N['coverage_flags']['register_not_recorded']} of {N['places']['all']} places have no recorded mobile status in the NTG register",
            "Regional counts only; star schema loads into DuckDB or Postgres",
        ], size=12, gap=0.9, dot_color=BLUE),
        t(9.09, 1.7, 3.5, 0.5, "Carriers (restricted)", size=20, bold=True, color=BLUE),
        *dotted_list(9.09, 2.35, 3.45, [
            "Relay register with hashed IDs and repair windows by month",
            f"Tenure flag on the {r['on_aboriginal_land_trust']} relays on Aboriginal Land Trust land",
            "Correction file: carriers fix the inferred network",
        ], size=12, gap=0.9, dot_color=BLUE),
        t(0.6, 6.3, 12.2, 0.5, "Offline: a Python app on a laptop, a 3 MB single-file page for phones and USB sticks, and a printable card.", size=14, bold=True),
        footer(),
    ], notes="Each audience gets its own output. Communities get a draft card to redesign with us. DCDD gets regional counts and KPIs "
             "with baselines. Carriers get the relay register, which stays restricted. Everything runs offline; the demo runs "
             "with Wi-Fi switched off. About 40 seconds."),

    dict(bg=DARK, els=[
        title("Ethics rules in the code, and what we ask for", color=WHITE),
        t(0.6, 1.5, 5.8, 0.5, "Rules the code enforces", size=18, bold=True, color=LIGHT),
        *dotted_list(0.6, 2.1, 5.8, [
            "We analyse the network and never rank communities.",
            "Public outputs show land council regions only, with small counts suppressed.",
            "A custodian the community chooses approves or withholds anything about it.",
            "Relay details go only to DCDD and carriers.",
            "Paid interpreters translate; no machine translation.",
        ], size=14, gap=0.72, color=WHITE),
        t(6.9, 1.5, 5.8, 0.5, "What we ask for", size=18, bold=True, color=LIGHT),
        *dotted_list(6.9, 2.1, 5.85, [
            "Make backhaul type, battery hours and outage logs a condition of NT co-investment.",
            f"Before each wet season, confirm battery hours at the {r['far_from_sealed_road']} single-path relays more than 10 km from a sealed road.",
            "After the 2026-27 wet season, test the card with one community that chooses to take part.",
        ], size=14, gap=1.05, color=WHITE),
        t(0.6, 6.15, 12, 0.6, "Public data only. No community has reviewed this work.", size=14, color=LIGHT),
        footer(dark=True),
    ], notes="The ethics rules are in the code and tested. We study equipment, never people, and nothing about a named community goes "
             "out without a custodian that community chooses. Three asks: make three fields a condition of co-investment, check "
             "battery hours at 43 relays before each wet season, and co-design the card with one community. Thank you. About 35 seconds."),
]


# ------------------------------------------------------------------ PowerPoint
def rgb(h):
    return RGBColor.from_string(h)


def add_rich(tf, text, size, bold, color, align):
    first = True
    for para in text.split("\n"):
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER}[align]
        p.line_spacing = 1.08
        for i, chunk in enumerate(re.split(r"\*\*(.+?)\*\*", para)):
            if not chunk:
                continue
            run = p.add_run(); run.text = chunk
            f = run.font; f.name = FONT; f.size = Pt(size); f.bold = bold or (i % 2 == 1); f.color.rgb = rgb(color)


def build_pptx(path):
    prs = Presentation(); prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    blank = prs.slide_layouts[6]
    for s in SLIDES:
        sl = prs.slides.add_slide(blank)
        sl.background.fill.solid(); sl.background.fill.fore_color.rgb = rgb(s["bg"])
        for kind, e in s["els"]:
            if kind == "text":
                tb = sl.shapes.add_textbox(Inches(e["x"]), Inches(e["y"]), Inches(e["w"]), Inches(e["h"]))
                tf = tb.text_frame; tf.word_wrap = True
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                tf.vertical_anchor = MSO_ANCHOR.TOP
                add_rich(tf, e["text"], e["size"], e["bold"], e["color"], e["align"])
            elif kind in ("box", "rect"):
                shp = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if kind == "box" else MSO_SHAPE.RECTANGLE,
                                          Inches(e["x"]), Inches(e["y"]), Inches(e["w"]), Inches(e["h"]))
                shp.fill.solid(); shp.fill.fore_color.rgb = rgb(e["fill"]); shp.line.fill.background(); shp.shadow.inherit = False
                if kind == "box":
                    shp.adjustments[0] = 0.06
            elif kind == "dot":
                shp = sl.shapes.add_shape(MSO_SHAPE.OVAL, Inches(e["x"]), Inches(e["y"]), Inches(e["d"]), Inches(e["d"]))
                shp.fill.solid(); shp.fill.fore_color.rgb = rgb(e["fill"]); shp.line.fill.background(); shp.shadow.inherit = False
            elif kind == "line":
                ln = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(e["x1"]), Inches(e["y1"]), Inches(e["x2"]), Inches(e["y2"]))
                ln.line.color.rgb = rgb(e["color"]); ln.line.width = Pt(e["width"])
            elif kind == "image":
                sl.shapes.add_picture(e["path"], Inches(e["x"]), Inches(e["y"]), width=Inches(e["w"]))
            elif kind == "table":
                rows = e["rows"]
                gt = sl.shapes.add_table(len(rows), len(rows[0]), Inches(e["x"]), Inches(e["y"]), Inches(e["w"]), Inches(0.5 * len(rows))).table
                for j, cw in enumerate(e["col_w"]):
                    gt.columns[j].width = Inches(cw)
                for i, row in enumerate(rows):
                    for j, val in enumerate(row):
                        cell = gt.cell(i, j); cell.fill.solid(); cell.fill.fore_color.rgb = rgb(DARK if i == 0 else (TINT if i % 2 else WHITE))
                        cell.margin_left = cell.margin_right = Inches(0.08); cell.margin_top = cell.margin_bottom = Inches(0.05)
                        tf = cell.text_frame; tf.word_wrap = True
                        add_rich(tf, val, e["size"], i == 0, WHITE if i == 0 else INK, "left")
        sl.notes_slide.notes_text_frame.text = s["notes"]
    prs.save(path)


# ------------------------------------------------------------------ PDF
def build_pdf(path):
    fd = "/System/Library/Fonts/Supplemental/"
    pdfmetrics.registerFont(TTFont("Arial", fd + "Arial.ttf")); pdfmetrics.registerFont(TTFont("Arial-Bold", fd + "Arial Bold.ttf"))
    pdfmetrics.registerFontFamily("Arial", normal="Arial", bold="Arial-Bold", italic="Arial", boldItalic="Arial-Bold")
    PT = 72.0
    cv = rl_canvas.Canvas(str(path), pagesize=(W * PT, H * PT), initialFontName="Arial")
    cv.setTitle("Underlink: Revealing the Hidden Dependencies Behind NT Connectivity"); cv.setAuthor(TEAM)

    def col(h):
        return colors.HexColor("#" + h)

    def para(e):
        txt = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", e["text"].replace("&", "&amp;").replace("<", "&lt;")).replace("\n", "<br/>")
        st = ParagraphStyle("s", fontName="Arial-Bold" if e["bold"] else "Arial", fontSize=e["size"], leading=e["size"] * 1.2,
                            textColor=col(e["color"]), alignment=TA_CENTER if e["align"] == "center" else TA_LEFT)
        p = Paragraph(txt, st)
        _, h = p.wrap(e["w"] * PT, e["h"] * PT * 3)
        p.drawOn(cv, e["x"] * PT, (H - e["y"]) * PT - h)

    for s in SLIDES:
        cv.setFillColor(col(s["bg"])); cv.rect(0, 0, W * PT, H * PT, stroke=0, fill=1)
        for kind, e in s["els"]:
            if kind == "text":
                para(e)
            elif kind in ("box", "rect"):
                cv.setFillColor(col(e["fill"]))
                y0 = (H - e["y"] - e["h"]) * PT
                if kind == "box":
                    cv.roundRect(e["x"] * PT, y0, e["w"] * PT, e["h"] * PT, 0.06 * min(e["w"], e["h"]) * PT, stroke=0, fill=1)
                else:
                    cv.rect(e["x"] * PT, y0, e["w"] * PT, e["h"] * PT, stroke=0, fill=1)
            elif kind == "dot":
                cv.setFillColor(col(e["fill"])); rr = e["d"] / 2
                cv.circle((e["x"] + rr) * PT, (H - e["y"] - rr) * PT, rr * PT, stroke=0, fill=1)
            elif kind == "line":
                cv.setStrokeColor(col(e["color"])); cv.setLineWidth(e["width"])
                cv.line(e["x1"] * PT, (H - e["y1"]) * PT, e["x2"] * PT, (H - e["y2"]) * PT)
            elif kind == "image":
                iw, ih = PILImage.open(e["path"]).size; hh = e["w"] * ih / iw
                cv.drawImage(e["path"], e["x"] * PT, (H - e["y"] - hh) * PT, e["w"] * PT, hh * PT, mask="auto")
            elif kind == "table":
                rows = e["rows"]
                st = ParagraphStyle("c", fontName="Arial", fontSize=e["size"], leading=e["size"] * 1.2, textColor=col(INK))
                sh = ParagraphStyle("h", parent=st, fontName="Arial-Bold", textColor=col(WHITE))
                data = [[Paragraph(v, sh if i == 0 else st) for v in row] for i, row in enumerate(rows)]
                tb = Table(data, colWidths=[w_ * PT for w_ in e["col_w"]])
                style = [("BACKGROUND", (0, 0), (-1, 0), col(DARK)), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                         ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                         ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5), ("FONTNAME", (0, 0), (-1, -1), "Arial")]
                style += [("BACKGROUND", (0, i), (-1, i), col(TINT if i % 2 else WHITE)) for i in range(1, len(rows))]
                tb.setStyle(TableStyle(style))
                _, th = tb.wrap(e["w"] * PT, H * PT)
                tb.drawOn(cv, e["x"] * PT, (H - e["y"]) * PT - th)
        cv.showPage()
    cv.save()


if __name__ == "__main__":
    build_pptx(HERE / f"{STEM}.pptx")
    build_pdf(HERE / f"{STEM}.pdf")
    print("wrote", HERE / f"{STEM}.pptx", "and", HERE / f"{STEM}.pdf")
